#!/usr/bin/env python3
"""
Rivera bilingual copy tool.

  extract  ->  read every data-en/data-vi (and placeholder) string from the site
               and write an editable spreadsheet (Rivera-copy-EN-VI.xlsx).
  apply    ->  read the edited spreadsheet back and write the English/Vietnamese
               values into the exact same attributes.

The mapping is by (page, ordinal): the Nth translatable element on a page maps to
ID "<page>-NNN". Both commands use the SAME parser, so the round-trip is exact as
long as the pages aren't hand-reordered in between.
"""
import sys, os, re, html

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "Rivera-copy-EN-VI.xlsx")

# reading order: the 3 main pages, then the 5 detail pages in collection order
PAGES = [
    "index.html", "projects.html", "about.html",
    "collins-wharf-aluna.html", "collins-wharf-ancora.html", "671-chapel-street.html",
    "380-melbourne.html", "aspire.html",
]

TOKEN = re.compile(r'<!--(?P<comment>.*?)-->|<(?P<tag>[a-zA-Z][^>]*?)>', re.DOTALL)

def _attr(tag, name):
    m = re.search(r'\b' + re.escape(name) + r'="([^"]*)"', tag)
    return m.group(1) if m else None

def _role(tagname, cls, field):
    if field == "placeholder":
        return "form placeholder"
    cls = cls or ""
    if "serif" in cls.split():
        return "italic accent (part of a heading)"
    if tagname in ("h1", "h2", "h3", "h4"):
        return "heading (%s)" % tagname.upper()
    if "eyebrow" in cls:
        return "eyebrow / small label"
    if "btn" in cls.split():
        return "button"
    if "lk" in cls.split() or "pd-back" in cls:
        return "text link"
    if "tag" in cls.split() or "st" in cls.split():
        return "tag / status chip"
    if "sub" in cls.split():
        return "hero subtitle"
    if "lead" in cls.split():
        return "lead paragraph"
    if "desc" in cls.split():
        return "card description"
    if "loc" in cls.split():
        return "location"
    if tagname == "small":
        return "spec label"
    if tagname == "b" or "pr" in cls.split() or "price" in cls.split():
        return "price / number"
    return "text"

def parse(page_path):
    """Yield dicts (n, field, en, vi, section, role, tagname) in document order."""
    txt = open(page_path, encoding="utf-8").read()
    section = ""
    n = 0
    for m in TOKEN.finditer(txt):
        if m.group("comment") is not None:
            c = " ".join(m.group("comment").split())
            if c:
                section = c[:48]
            continue
        tag = m.group(0)
        is_ph = "data-en-ph=" in tag
        if not is_ph and "data-en=" not in tag:
            continue
        n += 1
        tagname = re.match(r'<(\w+)', tag).group(1).lower()
        cls = _attr(tag, "class")
        if is_ph:
            en, vi = _attr(tag, "data-en-ph"), _attr(tag, "data-vi-ph")
            field = "placeholder"
        else:
            en, vi = _attr(tag, "data-en"), _attr(tag, "data-vi")
            field = "text"
        yield {
            "n": n, "field": field,
            "en": html.unescape(en or ""), "vi": html.unescape(vi or ""),
            "section": section, "role": _role(tagname, cls, field),
            "tagname": tagname,
        }

def _esc(v):
    return v.replace("&", "&amp;").replace('"', "&quot;")

# ---------------------------------------------------------------- extract
def extract():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = Workbook()

    info = wb.active
    info.title = "How to use"
    lines = [
        ("Rivera — website copy (English + Vietnamese)", True),
        ("", False),
        ("Edit ONLY the two yellow columns: 'English' and 'Vietnamese'.", False),
        ("Leave the ID, Page, Section, Role and Notes columns exactly as they are —", False),
        ("the ID is how each line is written back to the right spot on the site.", False),
        ("", False),
        ("• You can rewrite both languages freely. Keep them saying the same thing.", False),
        ("• Type normal text (apostrophes, &, accents are all fine) — escaping is handled.", False),
        ("• A few rows are flagged in Notes with 'keep the trailing space' — those are", False),
        ("  heading fragments that join to the next word; please keep that edge space.", False),
        ("• 'italic accent' rows are the emphasised serif phrase inside a heading.", False),
        ("• Blank a Vietnamese cell only if you truly want it identical to English.", False),
        ("", False),
        ("When done, save and send the file back — I'll write it into the pages.", False),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = info.cell(row=i, column=1, value=t)
        c.font = Font(name="Arial", size=12 if bold else 10, bold=bold)
    info.column_dimensions["A"].width = 96

    ws = wb.create_sheet("Copy")
    headers = ["ID", "Page", "Section", "Role", "Notes", "English", "Vietnamese"]
    ws.append(headers)
    lock_fill = PatternFill("solid", fgColor="EFEAE1")
    edit_fill = PatternFill("solid", fgColor="FFF6CC")
    hf = Font(name="Arial", size=10, bold=True)
    thin = Side(style="thin", color="D8D2C6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col)
        c.font = hf
        c.fill = edit_fill if h in ("English", "Vietnamese") else lock_fill
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border = border

    total = 0
    for page in PAGES:
        p = os.path.join(ROOT, page)
        if not os.path.exists(p):
            continue
        stem = page[:-5]
        for it in parse(p):
            note = []
            if it["en"] != it["en"].strip() or it["vi"] != it["vi"].strip():
                note.append("keep the leading/trailing space (joins with the adjacent word)")
            if it["role"].startswith("italic accent"):
                note.append("emphasised serif phrase")
            row = [
                "%s-%03d" % (stem, it["n"]), page, it["section"], it["role"],
                "; ".join(note), it["en"], it["vi"],
            ]
            ws.append(row)
            total += 1
            r = ws.max_row
            for col in range(1, 8):
                cell = ws.cell(row=r, column=col)
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=col in (3, 5, 6, 7))
                cell.border = border
                if col in (6, 7):
                    cell.fill = edit_fill

    widths = {"A": 22, "B": 16, "C": 30, "D": 26, "E": 34, "F": 58, "G": 58}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:G%d" % ws.max_row

    wb.save(XLSX)
    print("Wrote %s with %d copy rows across %d pages." % (XLSX, total, len(PAGES)))

# ---------------------------------------------------------------- apply
def apply(path):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["Copy"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    edits = {}  # id -> (en, vi)
    for r in rows:
        if not r or not r[0]:
            continue
        rid = str(r[0]).strip()
        en = "" if r[5] is None else str(r[5])
        vi = "" if r[6] is None else str(r[6])
        edits[rid] = (en, vi)

    changed_total = 0
    for page in PAGES:
        p = os.path.join(ROOT, page)
        if not os.path.exists(p):
            continue
        stem = page[:-5]
        txt = open(p, encoding="utf-8").read()
        # rebuild by walking the same tokens; replace attr values for translatable tags
        out = []
        last = 0
        n = 0
        changed = 0
        for m in TOKEN.finditer(txt):
            if m.group("comment") is not None:
                continue
            tag = m.group(0)
            is_ph = "data-en-ph=" in tag
            if not is_ph and "data-en=" not in tag:
                continue
            n += 1
            rid = "%s-%03d" % (stem, n)
            if rid not in edits:
                continue
            en, vi = edits[rid]
            en_a, vi_a = ("data-en-ph", "data-vi-ph") if is_ph else ("data-en", "data-vi")
            new = tag
            # only rewrite an attribute when its decoded value actually changed,
            # so untouched rows stay byte-identical (no spurious & re-escaping)
            cur_en = _attr(tag, en_a)
            if cur_en is None or html.unescape(cur_en) != en:
                new = re.sub(r'(\b%s=")[^"]*(")' % en_a, lambda mm: mm.group(1) + _esc(en) + mm.group(2), new, count=1)
            cur_vi = _attr(new, vi_a)
            if cur_vi is not None and html.unescape(cur_vi) != vi:
                new = re.sub(r'(\b%s=")[^"]*(")' % vi_a, lambda mm: mm.group(1) + _esc(vi) + mm.group(2), new, count=1)
            # keep the visible fallback text (between the tags) matching data-en,
            # but only for a clean text leaf: <tag ...>text</tag> with no child tags
            inner_repl = None
            if not is_ph and (cur_en is None or html.unescape(cur_en) != en):
                nxt = txt.find("<", m.end())
                tagname = re.match(r'<(\w+)', tag).group(1).lower()
                # clean leaf only: the very next tag must be this element's closing tag
                if nxt != -1 and txt[nxt:].lower().startswith("</" + tagname):
                    inner_repl = (m.end(), nxt, en.replace("&", "&amp;").replace("<", "&lt;"))
            if new != tag or inner_repl:
                out.append(txt[last:m.start()])
                out.append(new)
                if inner_repl:
                    out.append(inner_repl[2])
                    last = inner_repl[1]
                else:
                    last = m.end()
                changed += 1
        out.append(txt[last:])
        if changed:
            open(p, "w", encoding="utf-8").write("".join(out))
        changed_total += changed
        print("  %-26s %d strings updated" % (page, changed))
    print("Applied %d edits." % changed_total)

if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else "extract"
    if cmd == "extract":
        extract()
    elif cmd == "apply":
        apply(sys.argv[2] if len(sys.argv) > 2 else XLSX)
    else:
        print("usage: copy_tool.py [extract | apply <xlsx>]")
